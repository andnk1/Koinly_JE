"""
xlsx_report.py
==============
Builds the single 3-tab workbook the app hands back to the user:

    1. Reconciliation  -- beginning/ending balance control, Koinly vs Calculated
    2. JE by Month     -- one block per month, debit line(s) first within each
                          JE group, a thin rule under each JE group and a
                          thicker rule between months
    3. Total JE        -- the whole-period combined JE, plain (no separators,
                          natural line order -- matches JE_total.csv exactly)

QB_JE.csv is intentionally NOT part of this workbook -- it stays a separate
CSV download (that's the file that actually gets imported into QuickBooks,
and it needs to stay in QuickBooks' own import format).

Styled with the TaxPro UI kit's palette (navy / green / teal / grey) so it
reads as an extension of the web app itself.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from koinly_engine import JE_HEADER

# TaxPro palette (see static/style.css)
NAVY = "003366"
NAVY_LIGHT = "3C69AB"
GREEN = "446B67"
TEAL = "4E7874"
GREY = "9AA3B2"
CARDBG = "EFEFF7"
INK = "374151"
WHITE = "FFFFFF"

HEADER_FONT = Font(color=WHITE, bold=True, size=10, name="Calibri")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BODY_FONT = Font(color=INK, size=10.5, name="Calibri")
BOLD_FONT = Font(color=INK, size=10.5, bold=True, name="Calibri")
LABEL_FONT = Font(color=TEAL, bold=True, size=10, name="Calibri")
TITLE_FONT = Font(color=NAVY, bold=True, size=14, name="Calibri")
DIFF_OK_FONT = Font(color=GREEN, bold=True, size=10.5, name="Calibri")
DIFF_BAD_FONT = Font(color="8B0000", bold=True, size=10.5, name="Calibri")

THIN_BOTTOM = Border(bottom=Side(style="thin", color=GREY))
THICK_BOTTOM = Border(bottom=Side(style="medium", color=NAVY))
MONEY_FMT = "#,##0.00;[RED]-#,##0.00"


def _style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_je_rows(ws, start_row, records, thin_after_group=False, thick_between_months=False,
                   month_balances=None):
    """Writes JE_HEADER-shaped records starting at start_row. Optionally
    draws a thin rule after each (je_no, memo) group and a thicker rule
    between je_no (month) boundaries -- the thick rule wins where both would
    apply. Returns the next free row."""
    row = start_row
    n = len(records)
    i = 0
    while i < n:
        r = records[i]
        is_last_of_group = (i == n - 1) or (
            records[i + 1]["je_no"] != r["je_no"] or records[i + 1]["memo"] != r["memo"]
        )
        is_last_of_month = (i == n - 1) or (records[i + 1]["je_no"] != r["je_no"])

        vals = [r["je_date"], r["je_no"], r["acct"], r["gl"],
                r["debit"] or None, r["credit"] or None, r["memo"], r["name"]]
        if month_balances is not None:
            # Running balance shown once, on the last line of each month.
            vals.append(month_balances.get(r["je_no"]) if is_last_of_month else None)
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = BOLD_FONT if c == 9 else BODY_FONT
            if c in (5, 6, 9) and v is not None:
                cell.number_format = MONEY_FMT

        if thick_between_months and is_last_of_month:
            for c in range(1, len(vals) + 1):
                ws.cell(row=row, column=c).border = THICK_BOTTOM
        elif thin_after_group and is_last_of_group:
            for c in range(1, len(vals) + 1):
                ws.cell(row=row, column=c).border = THIN_BOTTOM

        row += 1
        i += 1
    return row


def _add_je_sheet(wb, title, records, thin_after_group, thick_between_months,
                  month_balances=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    header = list(JE_HEADER) + (["Balance"] if month_balances is not None else [])
    for c, h in enumerate(header, start=1):
        ws.cell(row=1, column=c, value=h)
    _style_header_row(ws, 1, len(header))
    ws.freeze_panes = "A2"
    next_row = _write_je_rows(ws, 2, records, thin_after_group, thick_between_months,
                              month_balances)
    _autosize(ws, [12, 12, 22, 22, 13, 13, 34, 12, 16])
    if not records:
        ws.cell(row=2, column=1, value="(no transactions in this period)").font = BODY_FONT
    return ws


def _add_reconciliation_sheet(wb, reconciliation):
    ws = wb.create_sheet("Reconciliation")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Reconciliation"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:C1")

    headers = ["", "Koinly", "Calculated"]
    header_row = 3
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    _style_header_row(ws, header_row, 3)

    r = reconciliation
    rows = [
        ("Beginning Balance", r["beginning"], r["beginning"]),
        ("D Digital Asset", None, r["debit"]),
        ("C Digital Assets", None, -r["credit"]),
        ("D Cash", None, r.get("cash_debit", 0.0)),
        ("C Cash", None, -r.get("cash_credit", 0.0)),
        ("Ending Balance", r["ending"], r["calculated_ending"]),
        ("", None, None),
        ("Difference", None, r["difference"]),
    ]
    row = header_row + 1
    for label, koinly_v, calc_v in rows:
        ws.cell(row=row, column=1, value=label).font = BOLD_FONT if label else BODY_FONT
        kc = ws.cell(row=row, column=2, value=koinly_v)
        cc = ws.cell(row=row, column=3, value=calc_v)
        for cell in (kc, cc):
            cell.font = BODY_FONT
            if cell.value is not None:
                cell.number_format = MONEY_FMT
        if label == "Difference":
            ok = abs(r["difference"]) < 0.01
            cc.font = DIFF_OK_FONT if ok else DIFF_BAD_FONT
            note = ws.cell(row=row, column=4,
                            value=("Ties out." if ok else "Does not tie to Koinly's ending balance -- check for missing transactions or a wrong beginning/ending balance."))
            note.font = Font(color=GREEN if ok else "8B0000", italic=True, size=9.5)
        row += 1

    note_row = row + 1
    ws.cell(row=note_row, column=1,
            value="Beginning/Ending Balance are optional -- a blank field is treated as 0.00 and never blocks the report.").font = \
        Font(color="6B7280", italic=True, size=9)

    _autosize(ws, [22, 16, 16, 60])
    return ws


def build_workbook(engine_result):
    """engine_result: the dict returned by koinly_engine.process(). Returns
    raw .xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    _add_reconciliation_sheet(wb, engine_result["reconciliation"])
    _add_je_sheet(wb, "JE by Month", engine_result["monthly_records_display"],
                  thin_after_group=True, thick_between_months=True,
                  month_balances=engine_result.get("monthly_balances", {}))
    _add_je_sheet(wb, "Total JE", engine_result["total_records"],
                  thin_after_group=False, thick_between_months=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
